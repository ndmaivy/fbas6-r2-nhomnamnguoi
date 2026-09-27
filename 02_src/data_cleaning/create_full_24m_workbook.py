"""Create one complete Excel workbook for the 24-month review dataset."""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
REVIEWS_PATH = ROOT / "01_data/processed/reviews_analysis_24m.csv"
CLASSIFIED_PATH = (
    ROOT / "01_data/processed/reviews_classified_24m_provisional_adjudicated.csv"
)
INCIDENTS_PATH = (
    ROOT / "01_data/processed/review_incidents_24m_provisional_adjudicated.csv"
)
OUTPUT_PATH = ROOT / "05_outputs/tables/TCInvest_736_reviews_24_months_full.xlsx"

KEY = ["source", "review_id"]
EXPECTED_REVIEWS = 736
EXPECTED_INCIDENTS = 796


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, encoding="utf-8-sig", dtype={"review_id": "string"})


def create_reviews_table() -> pd.DataFrame:
    reviews = read_csv(REVIEWS_PATH)
    classified = read_csv(CLASSIFIED_PATH)

    if len(reviews) != EXPECTED_REVIEWS or len(classified) != EXPECTED_REVIEWS:
        raise ValueError("The review inputs must each contain exactly 736 rows.")
    if reviews.duplicated(KEY).any() or classified.duplicated(KEY).any():
        raise ValueError("Duplicate (source, review_id) keys found in review inputs.")
    if set(map(tuple, reviews[KEY].to_numpy())) != set(
        map(tuple, classified[KEY].to_numpy())
    ):
        raise ValueError("The base and classified review keys do not match.")

    classification_columns = [
        column
        for column in classified.columns
        if column not in KEY
        and column not in {"review_date", "rating", "title", "review_text_raw"}
    ]
    merged = reviews.merge(
        classified[KEY + classification_columns],
        on=KEY,
        how="left",
        validate="one_to_one",
    )
    merged.insert(0, "stt", range(1, len(merged) + 1))
    merged = merged.sort_values(
        ["review_date", "source", "review_id"], ascending=[False, True, True]
    ).reset_index(drop=True)
    merged["stt"] = range(1, len(merged) + 1)
    return merged


def create_summary(reviews: pd.DataFrame, incidents: pd.DataFrame) -> pd.DataFrame:
    source_counts = reviews["source"].value_counts()
    values = [
        ("Tên bộ dữ liệu", "TCInvest public reviews - 24 tháng"),
        ("Khoảng thời gian phân tích", "22/09/2024 - 22/09/2026"),
        ("Ngày review sớm nhất thực tế", reviews["review_date"].min()),
        ("Ngày review muộn nhất thực tế", reviews["review_date"].max()),
        ("Tổng số review", len(reviews)),
        ("Google Play", int(source_counts.get("Google Play", 0))),
        ("App Store", int(source_counts.get("App Store", 0))),
        ("Khóa review duy nhất", reviews[KEY].drop_duplicates().shape[0]),
        ("Tổng số incident", len(incidents)),
        ("File review nguồn", REVIEWS_PATH.relative_to(ROOT).as_posix()),
        ("File phân loại nguồn", CLASSIFIED_PATH.relative_to(ROOT).as_posix()),
        ("File incident nguồn", INCIDENTS_PATH.relative_to(ROOT).as_posix()),
        (
            "Lưu ý phạm vi",
            "736 review records không đồng nghĩa với 736 khách hàng duy nhất.",
        ),
        (
            "Lưu ý phân loại",
            "Nhãn phân loại là PROVISIONAL; QA adjudication là self-review, "
            "chưa phải independent human validation.",
        ),
    ]
    return pd.DataFrame(values, columns=["Chỉ tiêu", "Giá trị"])


def create_dictionary(reviews: pd.DataFrame) -> pd.DataFrame:
    descriptions = {
        "stt": "Số thứ tự trong workbook, sắp xếp review mới nhất trước.",
        "source": "Nguồn review: Google Play hoặc App Store.",
        "review_id": "ID review tại nguồn; khóa duy nhất phải dùng cùng source.",
        "review_date": "Thời điểm review được ghi nhận tại nguồn.",
        "rating": "Điểm đánh giá từ 1 đến 5 sao.",
        "title": "Tiêu đề review nếu nguồn có cung cấp.",
        "review_text_raw": "Nội dung review nguyên bản.",
        "review_text_clean": "Nội dung được làm sạch nhẹ, không dịch hoặc diễn giải.",
        "analysis_window": "Cửa sổ thời gian dùng cho phân tích.",
        "incident_present": "Review có chứa ít nhất một incident được nhận diện.",
        "incident_count": "Số incident được nhận diện trong review.",
        "sentiment": "Nhãn cảm xúc ở cấp review.",
        "complaint_present": "Review có nội dung khiếu nại/phàn nàn.",
        "product_related": "Nội dung có liên quan đến sản phẩm.",
        "primary_pain_point_family": "Nhóm pain point chính.",
        "primary_pain_point_subtype": "Pain point subtype chính.",
        "primary_journey": "Customer journey chính.",
        "max_severity": "Mức độ nghiêm trọng cao nhất trong review.",
        "classification_confidence": "Độ tin cậy của nhãn phân loại.",
        "needs_human_review": "Cờ yêu cầu con người kiểm tra lại.",
        "classification_status": "Trạng thái của nhãn phân loại.",
    }
    return pd.DataFrame(
        {
            "Tên cột": reviews.columns,
            "Mô tả": [
                descriptions.get(column, "Trường kế thừa từ bộ dữ liệu review đã làm sạch.")
                for column in reviews.columns
            ],
        }
    )


def format_sheet(ws, freeze_panes: str, widths: dict[str, int] | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = freeze_panes
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32

    if widths:
        for column_name, width in widths.items():
            for cell in ws[1]:
                if cell.value == column_name:
                    ws.column_dimensions[get_column_letter(cell.column)].width = width
                    break


def main() -> None:
    reviews = create_reviews_table()
    incidents = read_csv(INCIDENTS_PATH)
    if len(incidents) != EXPECTED_INCIDENTS:
        raise ValueError("The incident input must contain exactly 796 rows.")
    if incidents["incident_id"].duplicated().any():
        raise ValueError("Duplicate incident_id values found.")
    if not set(map(tuple, incidents[KEY].to_numpy())).issubset(
        set(map(tuple, reviews[KEY].to_numpy()))
    ):
        raise ValueError("At least one incident does not map to a review.")

    summary = create_summary(reviews, incidents)
    dictionary = create_dictionary(reviews)
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="00_Tong_quan", index=False)
        reviews.to_excel(writer, sheet_name="01_736_reviews", index=False)
        incidents.to_excel(writer, sheet_name="02_796_incidents", index=False)
        dictionary.to_excel(writer, sheet_name="03_Tu_dien_du_lieu", index=False)

        format_sheet(
            writer.sheets["00_Tong_quan"],
            "A2",
            {"Chỉ tiêu": 34, "Giá trị": 90},
        )
        format_sheet(
            writer.sheets["01_736_reviews"],
            "G2",
            {
                "stt": 8,
                "source": 16,
                "review_id": 38,
                "review_date": 27,
                "rating": 10,
                "title": 35,
                "review_text_raw": 70,
                "review_text_clean": 70,
                "developer_reply": 60,
                "raw_source_file": 55,
                "primary_pain_point_family": 32,
                "primary_pain_point_subtype": 36,
                "primary_journey": 28,
                "classification_status": 22,
            },
        )
        format_sheet(
            writer.sheets["02_796_incidents"],
            "D2",
            {
                "incident_id": 32,
                "source": 16,
                "review_id": 38,
                "pain_point_family": 32,
                "pain_point_subtype": 36,
                "journey": 28,
                "evidence_span": 70,
                "classification_status": 22,
            },
        )
        format_sheet(
            writer.sheets["03_Tu_dien_du_lieu"],
            "A2",
            {"Tên cột": 36, "Mô tả": 90},
        )

        for sheet_name in ["01_736_reviews", "02_796_incidents"]:
            ws = writer.sheets[sheet_name]
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=False)

    print(f"Created: {OUTPUT_PATH}")
    print(f"Reviews: {len(reviews)}; columns: {len(reviews.columns)}")
    print(f"Incidents: {len(incidents)}")


if __name__ == "__main__":
    main()
