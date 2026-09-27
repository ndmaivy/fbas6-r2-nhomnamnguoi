#!/usr/bin/env python3
"""Rebuild the final 100-review manual taxonomy sample after label recovery."""

from __future__ import annotations

import hashlib
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
ANALYSIS_PATH = ROOT / "01_data/processed/reviews_analysis_24m.csv"
RECOVERED_PATH = ROOT / "05_outputs/tables/recovered_complete_manual_labels.csv"
INCOMPLETE_PATH = ROOT / "05_outputs/tables/TCInvest_manual_review_incomplete.xlsx"
BATCH1_PATH = ROOT / "05_outputs/tables/TCInvest_manual_review_batch01.xlsx"
BATCH2_PATH = ROOT / "05_outputs/tables/TCInvest_manual_review_batch02_labels_b.json"
SAMPLE_PATH = ROOT / "04_analysis/taxonomy/manual_review_sample_24m.csv"
SAMPLE_XLSX_PATH = ROOT / "04_analysis/taxonomy/manual_review_sample_24m.xlsx"
ARCHIVE_PATH = ROOT / "04_analysis/taxonomy/archive/manual_review_sample_24m_pre_recovery.csv"
SUMMARY_PATH = ROOT / "05_outputs/tables/manual_review_sample_24m_summary.csv"
OVERLAP_PATH = ROOT / "05_outputs/tables/manual_review_overlap_24m.csv"
REPORT_PATH = ROOT / "04_analysis/validation/manual_review_reuse_24m.md"

TARGET_SIZE = 100
RANDOM_SEED = 42
START = pd.Timestamp("2024-09-22T00:00:00Z")
END = pd.Timestamp("2026-09-22T23:59:59Z")
TARGET_RATING = {"LOW": 0.60, "MID": 0.15, "HIGH": 0.25}
TARGET_TIME = {"P1": 0.25, "P2": 0.25, "P3": 0.25, "P4": 0.25}
TARGET_LENGTH = {"VERY_SHORT": 0.10, "NORMAL": 0.65, "LONG": 0.25}

CORE_MANUAL_FIELDS = [
    "manual_validity", "sentiment", "complaint_present", "pain_point_primary",
    "pain_point_secondary_1", "pain_point_secondary_2", "journey_primary",
    "journey_secondary", "severity", "product_related", "feature_request",
    "evidence_span", "confidence", "reviewer_note", "needs_second_review",
]
ALL_MANUAL_FIELDS = [
    "label_status", "manual_validity", "invalid_reason", "sentiment",
    "complaint_present", "issue_summary", "pain_point_primary",
    "pain_point_secondary_1", "pain_point_secondary_2", "journey_primary",
    "journey_secondary", "severity", "product_related", "feature_request",
    "evidence_span", "confidence", "reviewer_note", "reviewer", "reviewed_at",
    "needs_second_review", "incidents_json", "label_origin", "completion_status",
    "incomplete_reason", "original_label_status", "recovery_source",
    "recovered_complete_basis", "source_review_date_value", "recovery_match_status",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def blank(value: object) -> bool:
    return pd.isna(value) or str(value).strip() == ""


def text(value: object) -> str:
    return "" if blank(value) else str(value)


def bool_true(value: object) -> bool:
    return text(value).strip().casefold() in {"true", "yes", "1"}


def largest_remainder(total: int, shares: dict[str, float]) -> dict[str, int]:
    raw = {key: total * value for key, value in shares.items()}
    result = {key: int(np.floor(value)) for key, value in raw.items()}
    order = sorted(shares, key=lambda key: (-(raw[key] - result[key]), key))
    for key in order[: total - sum(result.values())]:
        result[key] += 1
    return result


def complement_quota(targets: dict[str, int], existing: pd.Series, needed: int) -> dict[str, int]:
    deficits = {key: max(target - int(existing.get(key, 0)), 0) for key, target in targets.items()}
    deficit_total = sum(deficits.values())
    if deficit_total == needed:
        return deficits
    if deficit_total > needed:
        return largest_remainder(needed, {key: value / deficit_total for key, value in deficits.items()})
    remaining = needed - deficit_total
    base_total = sum(targets.values())
    extra = largest_remainder(remaining, {key: value / base_total for key, value in targets.items()})
    return {key: deficits[key] + extra[key] for key in targets}


def allocate_matrix(row_totals: dict[str, int], col_totals: dict[str, int]) -> dict[tuple[str, str], int]:
    total = sum(row_totals.values())
    if total != sum(col_totals.values()):
        raise RuntimeError("Allocation margins differ")
    raw = {(row, col): row_totals[row] * col_totals[col] / total for row in row_totals for col in col_totals}
    result = {key: int(np.floor(value)) for key, value in raw.items()}
    row_left = {row: row_totals[row] - sum(result[(row, col)] for col in col_totals) for row in row_totals}
    col_left = {col: col_totals[col] - sum(result[(row, col)] for row in row_totals) for col in col_totals}
    while sum(row_left.values()):
        choices = [
            (raw[(row, col)] - result[(row, col)], row, col)
            for row in row_totals for col in col_totals
            if row_left[row] and col_left[col]
        ]
        _, row, col = max(choices, key=lambda item: (item[0], item[1], item[2]))
        result[(row, col)] += 1
        row_left[row] -= 1
        col_left[col] -= 1
    return result


def rating_band(series: pd.Series) -> pd.Series:
    rating = pd.to_numeric(series, errors="coerce")
    return pd.Series(
        np.select([rating.le(2), rating.eq(3), rating.ge(4)], ["LOW", "MID", "HIGH"], default="UNKNOWN"),
        index=series.index,
    )


def time_period(series: pd.Series) -> pd.Series:
    dates = pd.to_datetime(series, utc=True, errors="coerce")
    return pd.cut(
        dates,
        bins=[
            pd.Timestamp("2024-09-22T00:00:00Z"), pd.Timestamp("2025-03-22T00:00:00Z"),
            pd.Timestamp("2025-09-22T00:00:00Z"), pd.Timestamp("2026-03-22T00:00:00Z"),
            pd.Timestamp("2026-09-23T00:00:00Z"),
        ],
        labels=["P1", "P2", "P3", "P4"],
        right=False,
    ).astype("string")


def length_band(frame: pd.DataFrame) -> pd.Series:
    length = pd.to_numeric(frame["review_char_length"], errors="coerce")
    short = frame["is_very_short"].map(bool_true)
    return pd.Series(
        np.select([short, length.ge(200)], ["VERY_SHORT", "LONG"], default="NORMAL"),
        index=frame.index,
    )


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for cells in sheet.columns:
            header = str(cells[0].value or "")
            letter = get_column_letter(cells[0].column)
            if header in {"review_text_raw", "review_text_clean", "issue_summary", "evidence_span", "reviewer_note", "incidents_json"}:
                width = 48
            elif header in {"review_id", "raw_source_file", "exact_duplicate_group_id"}:
                width = 38
            else:
                sample_widths = [len(str(cell.value)) for cell in cells[:101] if cell.value is not None]
                width = min(max([len(header), *sample_widths], default=10) + 2, 26)
            sheet.column_dimensions[letter].width = width
    workbook.save(path)


def main() -> None:
    protected = [ANALYSIS_PATH, RECOVERED_PATH, BATCH1_PATH, BATCH2_PATH, INCOMPLETE_PATH]
    hashes_before = {str(path.relative_to(ROOT)): sha256(path) for path in protected}

    analysis = pd.read_csv(ANALYSIS_PATH, dtype=str, keep_default_na=False)
    recovered = pd.read_csv(RECOVERED_PATH, dtype=str, keep_default_na=False)
    previous = pd.read_csv(SAMPLE_PATH, dtype=str, keep_default_na=False)
    incomplete = pd.read_excel(INCOMPLETE_PATH, sheet_name="Incomplete Reviews", dtype={"review_id": str})
    codebook = pd.read_excel(INCOMPLETE_PATH, sheet_name="Codebook")

    if len(analysis) != 736 or analysis["source"].value_counts().to_dict() != {"Google Play": 403, "App Store": 333}:
        raise RuntimeError("Analysis dataset does not match the validated 736-row scope")
    if analysis.duplicated(["source", "review_id"], keep=False).any():
        raise RuntimeError("Duplicate analysis key")
    analysis_dates = pd.to_datetime(analysis["review_date"], utc=True, errors="coerce")
    if analysis_dates.isna().any() or not analysis_dates.between(START, END, inclusive="both").all():
        raise RuntimeError("Analysis dates are invalid or outside the agreed window")

    second = incomplete[incomplete["label_status"].eq("NEEDS_SECOND_REVIEW")].copy()
    recovered["review_id"] = recovered["review_id"].astype(str)
    second["review_id"] = second["review_id"].astype(str)
    recovered_keys = set(zip(recovered["source"], recovered["review_id"]))
    second_keys = set(zip(second["source"], second["review_id"]))
    if recovered_keys & second_keys:
        raise RuntimeError("A key appears in both COMPLETE and second-review evidence")
    inspected_keys = recovered_keys | second_keys
    if len(recovered_keys) != len(recovered) or len(second_keys) != len(second):
        raise RuntimeError("Duplicate inspected keys")
    analysis_keys = set(zip(analysis["source"], analysis["review_id"]))
    if not inspected_keys.issubset(analysis_keys):
        raise RuntimeError("An inspected key is absent from the analysis dataset")

    existing_count = len(inspected_keys)
    new_needed = max(TARGET_SIZE - existing_count, 0)
    previous_keys = set(zip(previous["source"], previous["review_id"]))
    existing_in_previous = len(inspected_keys & previous_keys)
    existing_not_previous = len(inspected_keys - previous_keys)

    analysis_index = analysis.set_index(["source", "review_id"], drop=False)
    rows: list[dict[str, object]] = []
    for _, label in recovered.iterrows():
        key = (label["source"], label["review_id"])
        row = analysis_index.loc[key].to_dict()
        for field in ALL_MANUAL_FIELDS:
            row[field] = label[field] if field in label.index else ""
        row["label_status"] = "REVIEWED"
        row["manual_reuse_status"] = "REUSED_COMPLETE"
        rows.append(row)
    for _, label in second.iterrows():
        key = (label["source"], label["review_id"])
        row = analysis_index.loc[key].to_dict()
        for field in ALL_MANUAL_FIELDS:
            row[field] = label[field] if field in label.index else ""
        row["label_status"] = "NEEDS_SECOND_REVIEW"
        row["manual_reuse_status"] = "REUSED_NEEDS_SECOND_REVIEW"
        rows.append(row)
    existing = pd.DataFrame(rows)
    existing["rating_band"] = rating_band(existing["rating"])
    existing["time_period_24m"] = time_period(existing["review_date"])
    existing["text_length_band"] = length_band(existing)

    source_targets = largest_remainder(TARGET_SIZE, (analysis["source"].value_counts() / len(analysis)).to_dict())
    rating_targets = largest_remainder(TARGET_SIZE, TARGET_RATING)
    time_targets = largest_remainder(TARGET_SIZE, TARGET_TIME)
    length_targets = largest_remainder(TARGET_SIZE, TARGET_LENGTH)
    source_new = complement_quota(source_targets, existing["source"].value_counts(), new_needed)
    rating_new = complement_quota(rating_targets, existing["rating_band"].value_counts(), new_needed)
    time_new = complement_quota(time_targets, existing["time_period_24m"].value_counts(), new_needed)
    length_new = complement_quota(length_targets, existing["text_length_band"].value_counts(), new_needed)

    previous_new_keys = set(zip(
        previous.loc[previous["manual_reuse_status"].eq("NEW_SAMPLE_UNREVIEWED"), "source"],
        previous.loc[previous["manual_reuse_status"].eq("NEW_SAMPLE_UNREVIEWED"), "review_id"],
    )) - inspected_keys
    candidate_keys = pd.Series(list(zip(analysis["source"], analysis["review_id"])), index=analysis.index)
    candidates = analysis[~candidate_keys.isin(inspected_keys)].copy()
    candidates["_rating_band"] = rating_band(candidates["rating"])
    candidates["_time_period"] = time_period(candidates["review_date"])
    candidates["_length_band"] = length_band(candidates)
    candidates["_previous_candidate"] = [key in previous_new_keys for key in zip(candidates["source"], candidates["review_id"])]
    rng = np.random.default_rng(RANDOM_SEED)
    candidates["_tie"] = rng.random(len(candidates))

    matrix = allocate_matrix(source_new, rating_new) if new_needed else {}
    slots = [key for key, count in matrix.items() for _ in range(count)]
    rng.shuffle(slots)
    selected: list[int] = []
    time_selected = {key: 0 for key in time_new}
    length_selected = {key: 0 for key in length_new}
    for source_name, band in slots:
        pool = candidates[
            candidates["source"].eq(source_name)
            & candidates["_rating_band"].eq(band)
            & ~candidates.index.isin(selected)
        ].copy()
        if pool.empty:
            raise RuntimeError(f"No candidate for source={source_name}, rating={band}")
        pool["_score"] = (
            pool["_time_period"].map(lambda value: max(time_new.get(value, 0) - time_selected.get(value, 0), 0) / max(time_new.get(value, 1), 1)) * 5
            + pool["_length_band"].map(lambda value: max(length_new.get(value, 0) - length_selected.get(value, 0), 0) / max(length_new.get(value, 1), 1)) * 4
            + pool["_previous_candidate"].astype(int) * 0.25
            + pool["_tie"] * 1e-6
        )
        index = int(pool["_score"].idxmax())
        selected.append(index)
        chosen = candidates.loc[index]
        time_selected[chosen["_time_period"]] = time_selected.get(chosen["_time_period"], 0) + 1
        length_selected[chosen["_length_band"]] = length_selected.get(chosen["_length_band"], 0) + 1

    new = candidates.loc[selected, analysis.columns].copy()
    for field in ALL_MANUAL_FIELDS:
        new[field] = ""
    new["label_status"] = "UNREVIEWED"
    new["manual_reuse_status"] = "NEW_SAMPLE_UNREVIEWED"
    new["rating_band"] = rating_band(new["rating"])
    new["time_period_24m"] = time_period(new["review_date"])
    new["text_length_band"] = length_band(new)
    new["sampling_seed"] = RANDOM_SEED
    new["was_in_pre_recovery_sample"] = [key in previous_new_keys for key in zip(new["source"], new["review_id"])]

    existing["sampling_seed"] = ""
    existing["was_in_pre_recovery_sample"] = [key in previous_keys for key in zip(existing["source"], existing["review_id"])]
    final = pd.concat([existing, new], ignore_index=True, sort=False)
    order = {"REUSED_COMPLETE": 0, "REUSED_NEEDS_SECOND_REVIEW": 1, "NEW_SAMPLE_UNREVIEWED": 2}
    final["_order"] = final["manual_reuse_status"].map(order)
    final = final.sort_values(["_order", "review_date", "source", "review_id"]).drop(columns="_order").reset_index(drop=True)

    # Validate labels before writing.
    final_keys = set(zip(final["source"], final["review_id"]))
    if len(final) != max(TARGET_SIZE, existing_count) or len(final_keys) != len(final):
        raise RuntimeError("Final sample size/key validation failed")
    if not final_keys.issubset(analysis_keys):
        raise RuntimeError("Final sample contains a key outside analysis data")
    final_dates = pd.to_datetime(final["review_date"], utc=True, errors="coerce")
    if final_dates.isna().any() or not final_dates.between(START, END, inclusive="both").all():
        raise RuntimeError("Final sample contains an invalid/out-of-window date")

    new_mask = final["manual_reuse_status"].eq("NEW_SAMPLE_UNREVIEWED")
    new_blank = final.loc[new_mask, [field for field in ALL_MANUAL_FIELDS if field != "label_status"]].apply(
        lambda column: column.map(blank)
    ).all().all() and final.loc[new_mask, "label_status"].eq("UNREVIEWED").all()
    if not new_blank:
        raise RuntimeError("A new row contains a manual label")

    complete_preserved = True
    for _, source_row in recovered.iterrows():
        result = final[final["source"].eq(source_row["source"]) & final["review_id"].eq(source_row["review_id"])].iloc[0]
        for field in [field for field in ALL_MANUAL_FIELDS if field in recovered.columns and field != "label_status"]:
            if text(result[field]) != text(source_row[field]):
                complete_preserved = False
                break
    if not complete_preserved:
        raise RuntimeError("A recovered COMPLETE label changed")

    second_preserved = True
    for _, source_row in second.iterrows():
        result = final[final["source"].eq(source_row["source"]) & final["review_id"].eq(source_row["review_id"])].iloc[0]
        for field in [field for field in ALL_MANUAL_FIELDS if field in second.columns]:
            if text(result[field]).strip() != text(source_row[field]).strip():
                second_preserved = False
                break
    if not second_preserved:
        raise RuntimeError("A second-review label changed")

    # Taxonomy coverage uses only the 63 COMPLETE labels.
    complete = final[final["manual_reuse_status"].eq("REUSED_COMPLETE")]
    pain_counts = complete["pain_point_primary"].map(text).value_counts()
    journey_counts = complete["journey_primary"].map(text).value_counts()
    severity_counts = complete["severity"].map(text).value_counts()
    sentiment_counts = complete["sentiment"].map(text).value_counts()
    complaint_counts = complete["complaint_present"].map(text).value_counts()
    product_counts = complete["product_related"].map(text).value_counts()
    feature_counts = complete["feature_request"].map(text).value_counts()
    secondary_count = int(complete[["pain_point_secondary_1", "pain_point_secondary_2"]].apply(
        lambda row: any(text(value).strip().upper() not in {"", "NONE"} for value in row), axis=1
    ).sum())
    multi_count = secondary_count
    unclear_count = int(complete[["pain_point_primary", "pain_point_secondary_1", "pain_point_secondary_2", "journey_primary", "journey_secondary"]].apply(
        lambda row: any(text(value).strip().upper() == "UNCLEAR" for value in row), axis=1
    ).sum())
    none_count = int(complete["pain_point_primary"].map(text).str.upper().eq("NONE").sum())
    other_count = int(complete[["pain_point_primary", "journey_primary"]].apply(
        lambda row: any(text(value).strip().upper() == "OTHER" for value in row), axis=1
    ).sum())
    feature_request_count = int(complete["feature_request"].map(text).str.upper().eq("YES").sum())

    pain_codebook = [text(value).strip() for value in codebook["Pain point values"] if text(value).strip()]
    journey_codebook = [text(value).strip() for value in codebook["Journey values"] if text(value).strip()]
    pain_represented = [value for value in pain_codebook if value in pain_counts.index]
    journey_represented = [value for value in journey_codebook if value in journey_counts.index]
    pain_zero = [value for value in pain_codebook if value not in pain_counts.index]
    journey_zero = [value for value in journey_codebook if value not in journey_counts.index]

    # Keep the old generated sample as immutable history before overwriting it.
    ARCHIVE_PATH.parent.mkdir(parents=True, exist_ok=True)
    if not ARCHIVE_PATH.exists():
        shutil.copy2(SAMPLE_PATH, ARCHIVE_PATH)

    SAMPLE_PATH.parent.mkdir(parents=True, exist_ok=True)
    final.to_csv(SAMPLE_PATH, index=False)

    # Human-friendly workbook: requested fields first, then remaining metadata.
    priority = [
        "source", "review_id", "review_date", "rating", "review_text_raw",
        "manual_reuse_status", "label_status", "manual_validity", "sentiment",
        "complaint_present", "pain_point_primary", "pain_point_secondary_1",
        "pain_point_secondary_2", "journey_primary", "journey_secondary", "severity",
        "product_related", "feature_request", "evidence_span", "confidence",
        "reviewer_note", "needs_second_review",
    ]
    workbook_columns = priority + [column for column in final.columns if column not in priority]

    summary_dimensions = {
        "source": final["source"],
        "rating_band": final["rating_band"],
        "manual_reuse_status": final["manual_reuse_status"],
        "review_status": final["label_status"],
        "very_short_status": final["is_very_short"].map(lambda value: "VERY_SHORT" if bool_true(value) else "NOT_VERY_SHORT"),
        "time_period_24m": final["time_period_24m"],
        "text_length_band": final["text_length_band"],
        "pre_recovery_sample": final["was_in_pre_recovery_sample"].map(lambda value: "YES" if bool_true(value) else "NO"),
    }
    summary_rows = []
    for dimension, values in summary_dimensions.items():
        for category, count in values.value_counts(dropna=False).items():
            summary_rows.append({"dimension": dimension, "category": category, "count": int(count), "percentage": round(100 * count / len(final), 2)})
    summary = pd.DataFrame(summary_rows)
    summary.to_csv(SUMMARY_PATH, index=False)

    overlap = final[final["manual_reuse_status"].ne("NEW_SAMPLE_UNREVIEWED")].copy()
    overlap["old_label_status"] = overlap["label_status"]
    overlap["in_24m"] = True
    overlap_priority = ["source", "review_id", "review_date", "rating", "review_text_raw", "old_label_status", "needs_second_review", "in_24m", "manual_reuse_status"]
    overlap = overlap[overlap_priority + [column for column in overlap.columns if column not in overlap_priority]]
    overlap.to_csv(OVERLAP_PATH, index=False)

    workload = pd.DataFrame([
        ["COMPLETE — no action", int(final["manual_reuse_status"].eq("REUSED_COMPLETE").sum())],
        ["NEEDS_SECOND_REVIEW — human resolution", int(final["manual_reuse_status"].eq("REUSED_NEEDS_SECOND_REVIEW").sum())],
        ["NEW_SAMPLE_UNREVIEWED — human labeling", int(final["manual_reuse_status"].eq("NEW_SAMPLE_UNREVIEWED").sum())],
    ], columns=["Workload", "Reviews"])
    with pd.ExcelWriter(SAMPLE_XLSX_PATH, engine="openpyxl") as writer:
        final[workbook_columns].to_excel(writer, sheet_name="Manual Review", index=False)
        workload.to_excel(writer, sheet_name="Review Workload", index=False)
        summary.to_excel(writer, sheet_name="Sample Summary", index=False)
        codebook.to_excel(writer, sheet_name="Codebook", index=False)
    style_workbook(SAMPLE_XLSX_PATH)

    selected_previous = int(new["was_in_pre_recovery_sample"].sum())
    report = f"""# Manual Review Reuse — 24-Month Final Sample

## Recovery Status

- COMPLETE labels recovered: **{len(recovered)}**
- NEEDS_SECOND_REVIEW labels retained: **{len(second)}**
- Unique inspected evidence: **{existing_count}**
- Inspected keys already present in the pre-recovery sample: **{existing_in_previous}**
- Inspected keys absent from the pre-recovery sample: **{existing_not_previous}**

All inspected evidence was matched by `(source, review_id)` and retained. COMPLETE labels came only from `recovered_complete_manual_labels.csv`; second-review labels came from `TCInvest_manual_review_incomplete.xlsx`. No label was inferred or resolved.

## Rebuild and Sampling Method

- Final target: **{TARGET_SIZE} unique reviews**
- New reviews required and selected: **{new_needed}**
- New rows reused from the pre-recovery sample: **{selected_previous}**
- Additional new rows drawn from the analysis dataset: **{new_needed - selected_previous}**
- Random seed: **{RANDOM_SEED}**

The 70 inspected rows were retained regardless of imbalance. Their starting distribution was Google Play/App Store = {existing['source'].value_counts().to_dict()}, rating bands = {existing['rating_band'].value_counts().to_dict()}, time periods = {existing['time_period_24m'].value_counts().sort_index().to_dict()}, and text-length bands = {existing['text_length_band'].value_counts().to_dict()}.

New selection used dynamically calculated deficits against these diversity targets:

- Source: approximate 24-month dataset mix where feasible
- Rating: LOW 60%, MID 15%, HIGH 25%
- Time: four equal six-month periods
- Length: VERY_SHORT 10%, NORMAL 65%, LONG 25% (`LONG >= 200` characters)

Because existing evidence was heavily Google Play/P4, source and time targets were not achievable without dropping inspected rows. New rows therefore prioritize App Store and earlier periods. Previously sampled unreviewed rows receive a deterministic preference when they also fill diversity deficits.

## Final Sample Validation

- Rows and unique keys: **{len(final)} / {len(final_keys)}**
- COMPLETE: **{int(final['manual_reuse_status'].eq('REUSED_COMPLETE').sum())}**
- NEEDS_SECOND_REVIEW: **{int(final['manual_reuse_status'].eq('REUSED_NEEDS_SECOND_REVIEW').sum())}**
- NEW_SAMPLE_UNREVIEWED: **{int(final['manual_reuse_status'].eq('NEW_SAMPLE_UNREVIEWED').sum())}**
- Sources: {final['source'].value_counts().to_dict()}
- Rating bands: {final['rating_band'].value_counts().to_dict()}
- Time periods: {final['time_period_24m'].value_counts().sort_index().to_dict()}
- Text-length bands: {final['text_length_band'].value_counts().to_dict()}
- All keys belong to the 736-row dataset: **Yes**
- COMPLETE labels preserved exactly: **{'Yes' if complete_preserved else 'No'}**
- Second-review labels preserved: **{'Yes' if second_preserved else 'No'}**
- New manual fields blank: **{'Yes' if new_blank else 'No'}**

## Taxonomy Coverage from 63 COMPLETE Labels

- Pain-point primary: {pain_counts.to_dict()}
- Journey primary: {journey_counts.to_dict()}
- Severity: {severity_counts.to_dict()}
- Sentiment: {sentiment_counts.to_dict()}
- Complaint present: {complaint_counts.to_dict()}
- Product related: {product_counts.to_dict()}
- Feature request: {feature_counts.to_dict()}
- Reviews with secondary pain points: **{secondary_count}**
- Multi-pain-point reviews: **{multi_count}**
- Reviews containing UNCLEAR: **{unclear_count}**
- Primary pain point NONE: **{none_count}**
- Primary pain point or journey OTHER: **{other_count}**
- Feature-request reviews: **{feature_request_count}**

Represented pain-point categories: {pain_represented}

Pain-point categories with zero evidence: {pain_zero}

Represented journey categories: {journey_represented}

Journey categories with zero evidence: {journey_zero}

Zero evidence does not imply that a Codebook category should be removed.

## Human Review Workload

- COMPLETE — no action: **{len(recovered)}**
- NEEDS_SECOND_REVIEW — human resolution: **{len(second)}**
- NEW_SAMPLE_UNREVIEWED — human labeling: **{new_needed}**

## Limitations

- The sample is intentionally optimized for taxonomy discovery/validation, not prevalence estimation.
- The 70 historical inspected rows are source/time imbalanced and were not dropped.
- No taxonomy refinement, classification, second-review resolution, pain-point ranking or business inference was performed.
"""
    REPORT_PATH.write_text(report, encoding="utf-8")

    # Re-open outputs and prove protected sources stayed unchanged.
    check = pd.read_csv(SAMPLE_PATH, dtype=str, keep_default_na=False)
    if len(check) != len(final) or check.duplicated(["source", "review_id"], keep=False).any():
        raise RuntimeError("Written CSV failed final validation")
    hashes_after = {str(path.relative_to(ROOT)): sha256(path) for path in protected}
    if hashes_before != hashes_after:
        raise RuntimeError("A protected source file changed")

    print("FINAL MANUAL SAMPLE REBUILD RESULT")
    print(f"Recovered/second/existing: {len(recovered)}/{len(second)}/{existing_count}")
    print(f"Existing in/not previous: {existing_in_previous}/{existing_not_previous}")
    print(f"New needed/selected: {new_needed}/{len(new)}; prior sample reused: {selected_previous}")
    print(f"Final: {len(final)}; status={final['manual_reuse_status'].value_counts().to_dict()}")
    print(f"Sources: {final['source'].value_counts().to_dict()}")
    print(f"Rating bands: {final['rating_band'].value_counts().to_dict()}")
    print(f"Very short: {int(final['is_very_short'].map(bool_true).sum())}")
    print(f"Time: {final_dates.min().isoformat()} to {final_dates.max().isoformat()}, months={final_dates.dt.strftime('%Y-%m').nunique()}, periods={final['time_period_24m'].value_counts().sort_index().to_dict()}")
    print(f"Length: {final['text_length_band'].value_counts().to_dict()}")
    print(f"Pain represented/zero: {pain_represented}/{pain_zero}")
    print(f"Journey represented/zero: {journey_represented}/{journey_zero}")
    print(f"Severity: {severity_counts.to_dict()}")
    print(f"Sentiment: {sentiment_counts.to_dict()}")
    print(f"Feature/secondary: {feature_request_count}/{secondary_count}")
    print(f"Validation duplicate=0, membership=YES, complete_preserved={'YES' if complete_preserved else 'NO'}, second_preserved={'YES' if second_preserved else 'NO'}, new_blank={'YES' if new_blank else 'NO'}")
    print("Source files unchanged: YES")


if __name__ == "__main__":
    main()
